"""Shared Word styles for the AICC final-year report."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = RGBColor(0x0B, 0x3D, 0x66)
BLACK = RGBColor(0x11, 0x11, 0x11)


def _set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _set_paragraph_font(paragraph, size=12, bold=False, italic=False):
    for run in paragraph.runs:
        _set_run_font(run, size=size, bold=bold, italic=italic)


def shade_cell(cell, fill):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "8AA0B4")
        tc_borders.append(element)
    tc_pr.append(tc_borders)


def prevent_row_split(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def add_page_number(paragraph):
    run = paragraph.add_run()
    _set_run_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


def configure_section(section, start=None, fmt=None):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = False
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.clear()
    run = paragraph.add_run("AICC  |  Lahore Garrison University  |  ")
    _set_run_font(run, size=10, color=NAVY)
    add_page_number(paragraph)
    if start is not None or fmt is not None:
        sect_pr = section._sectPr
        pg_num = sect_pr.find(qn("w:pgNumType"))
        if pg_num is None:
            pg_num = OxmlElement("w:pgNumType")
            sect_pr.append(pg_num)
        if fmt:
            pg_num.set(qn("w:fmt"), fmt)
        if start is not None:
            pg_num.set(qn("w:start"), str(start))


def new_document():
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(8)
    pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    sizes = {1: 16, 2: 14, 3: 12}
    for level, size in sizes.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.font.italic = False
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        sp = style.paragraph_format
        sp.line_spacing = 1.15
        sp.space_before = Pt(12 if level > 1 else 0)
        sp.space_after = Pt(8)
        sp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        sp.keep_with_next = True
    configure_section(document.sections[0], start=1, fmt="lowerRoman")
    return document


def page_break(document):
    document.add_page_break()


def center(document, text, size=12, bold=False, italic=False, before=0, after=6, color=BLACK):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return paragraph


def body(document, text, first_line=True):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(8)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(1.0)
    run = paragraph.add_run(text)
    _set_run_font(run)
    return paragraph


def bullet(document, text, level=0):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.left_indent = Cm(1.25 + level * 0.75)
    paragraph.paragraph_format.first_line_indent = Cm(-0.4)
    run = paragraph.add_run("•  " + text)
    _set_run_font(run, size=12)
    return paragraph


def heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        _set_run_font(run, size={1: 16, 2: 14, 3: 12}[level], bold=True, color=NAVY)
    return paragraph


def caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    _set_run_font(run, size=11, bold=True, italic=False, color=NAVY)
    return paragraph


def table_caption(document, text):
    paragraph = caption(document, text)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_table(document, headers, rows):
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(header)
        _set_run_font(run, size=10, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(cell, "0B3D66")
        set_cell_border(cell)
    prevent_row_split(table.rows[0])
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.rows[r_index + 1].cells[c_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            _set_run_font(run, size=10)
            if r_index % 2 == 1:
                shade_cell(cell, "F4F7FA")
            set_cell_border(cell)
        prevent_row_split(table.rows[r_index + 1])
    document.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def add_figure(document, path, width=6.3):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run()
    run.add_picture(path, width=Inches(width))
    return paragraph


def toc_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose Update Field after opening in Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run2 = paragraph.add_run()
    run2._r.append(placeholder)
    run3 = paragraph.add_run()
    run3._r.append(end)
    _set_run_font(run, size=12, italic=True)


def enable_update_fields(document):
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
