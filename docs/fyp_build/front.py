"""Front matter matching the LGU OBE final-documentation template."""

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from styles import BLACK, NAVY, body, bullet, center, heading, page_break, toc_field
from styles import _set_run_font


TITLE = "AICC: An AI Campus Companion for Academic Assistance and Campus Navigation"
SHORT = "AICC: An AI Campus Companion"


def _line(document, text, size=12, bold=False, italic=False, before=0, after=4, color=BLACK):
    return center(document, text, size=size, bold=bold, italic=italic, before=before, after=after, color=color)


def add_front(document):
    _line(document, "LAHORE GARRISON UNIVERSITY", 18, True, before=36, after=4, color=NAVY)
    _line(document, "Department of Computer Science", 14, True, after=18, color=NAVY)
    _line(document, "Final Year Project", 13, True, before=12, after=16)
    _line(document, TITLE, 16, True, before=12, after=18, color=NAVY)
    _line(document, "Submitted by", 12, True, before=18, after=8)
    _line(document, "Abdul  (BSCS–________–____)", 13, True, after=4)
    _line(document, "________________  (BSCS–________–____)", 12, after=4)
    _line(document, "________________  (BSCS–________–____)", 12, after=4)
    _line(document, "________________  (BSCS–________–____)", 12, after=12)
    _line(document, "Session  2023–2026", 12, True, before=8, after=12)
    _line(document, "Supervised by", 12, True, before=8, after=6)
    _line(document, "_______________________________", 12, after=2)
    _line(document, "Teacher Name", 12, italic=True, after=16)
    _line(document, "Department of Computer Science", 12, True, before=10, after=2)
    _line(document, "Lahore Garrison University", 12, True, after=2)
    _line(document, "Lahore", 12, True, after=18)
    body(
        document,
        "Complete the underlined blanks before binding: full name, roll number, section, any co-authors, and the internal supervisor. Those details are not stored in the project repository, so they are left blank here rather than guessed.",
        first_line=False,
    )

    page_break(document)
    _line(document, TITLE, 16, True, before=12, after=16, color=NAVY)
    _line(document, "A project submitted to the", 12, before=8)
    _line(document, "Department of Computer Science", 13, True, after=8)
    _line(document, "In", 12, before=6, after=6)
    _line(document, "Partial Fulfillment of the Requirements for the", 12, after=4)
    _line(document, "Bachelor’s Degree in Computer Science", 13, True, after=14)
    _line(document, "By", 12, before=8, after=6)
    _line(document, "Abdul", 13, True, after=2)
    _line(document, "________________", 12, after=2)
    _line(document, "________________", 12, after=2)
    _line(document, "________________", 12, after=16)
    _line(document, "Internal Supervisor", 12, True, before=10, after=8)
    _line(document, "_______________________________          _______________________", 12, after=2)
    _line(document, "Teacher Name                                          Designation", 11, italic=True, after=2)
    _line(document, "Department of Computer Science", 12, after=14)
    _line(document, "External Examiner", 12, True, before=8, after=6)
    _line(document, "_______________________________", 12, after=16)
    _line(document, "Chairperson", 12, True, before=8, after=8)
    _line(document, "Dr. Erum Mehmood                        _______________________", 12, after=2)
    _line(document, "Assistant Professor", 12, italic=True, after=2)
    _line(document, "Department of Computer Science", 12, after=4)

    page_break(document)
    heading(document, "Copyright", 1)
    body(
        document,
        "This is to certify that the project titled “AICC: An AI Campus Companion for Academic Assistance and Campus Navigation” is the genuine work carried out by Abdul and the group members named on the declaration page, students of BS Computer Science in the Department of Computer Science, Lahore Garrison University, Lahore, during the academic session 2023–2026, in partial fulfilment of the requirements for the award of the degree of Bachelor of Science in Computer Science, and that the project has not formed the basis for the award previously of any other degree, diploma, fellowship, or any other similar title.",
    )
    _line(document, "Abdul  ______________________________", 12, before=18, after=8)
    _line(document, "________________  ______________________________", 12, after=8)
    _line(document, "________________  ______________________________", 12, after=8)
    _line(document, "________________  ______________________________", 12, after=8)

    page_break(document)
    heading(document, "Declaration", 1)
    body(
        document,
        "This is to declare that the project entitled “AICC: An AI Campus Companion for Academic Assistance and Campus Navigation” is an original work done by the undersigned, in partial fulfilment of the requirements for the degree “Bachelor of Science in Computer Science” at the Department of Computer Science, Lahore Garrison University, Lahore.",
    )
    body(
        document,
        "All the analysis, design, and system development described in this report have been accomplished by the undersigned. This project has not been submitted to any other college or university. Handbook answers used by the assistant are drawn from institutional student-information material and are not presented as a new statement of university policy. Where a rule and a live office notice disagree, the office notice remains authoritative.",
    )
    _line(document, "Group Members", 13, True, before=12, after=10, color=NAVY)
    _line(document, "Abdul                          ________________", 12, after=6)
    _line(document, "________________          ________________", 12, after=6)
    _line(document, "________________          ________________", 12, after=6)
    _line(document, "________________          ________________", 12, after=12)
    _line(document, "Supervisor", 13, True, before=8, after=8, color=NAVY)
    _line(document, "________________          ________________", 12, after=10)
    _line(document, "Date: ________________", 12, after=6)

    page_break(document)
    heading(document, "Dedication", 1)
    body(
        document,
        "This work is dedicated to our families, whose patience covered the long evenings of building, testing, and walking the campus with a phone in hand, and to the students who lose time looking for a room, a teacher, or a document that should have been one question away.",
    )
    body(
        document,
        "It is also dedicated to the teachers who insist that a useful system must be explainable. A campus assistant that invents a room number or a fee is not a convenience. It is a new problem. This project tries to be the other kind.",
    )

    page_break(document)
    heading(document, "Acknowledgements", 1)
    body(
        document,
        "We thank the Department of Computer Science at Lahore Garrison University for the chance to treat a daily campus difficulty as a final-year engineering problem. We thank our internal supervisor for review of the scope, for insisting that personal academic records stay separate from generated answers, and for the time given to design discussions. The supervisor’s name is to be entered on the title page.",
    )
    body(
        document,
        "We thank the chairperson and the faculty who reviewed intermediate demonstrations. We thank classmates who tried the chat in English and in Roman Urdu and told us when a place name failed. We thank the authors of the open tools used in the implementation, including Flutter, FastAPI, Firebase, and the map libraries, and the authors of the classical graph-search and retrieval papers cited in this report.",
    )
    body(
        document,
        "Any remaining error in a route, a coordinate, or a handbook paraphrase is ours. The university’s own offices remain the authority for attendance, fees, examinations, and conduct.",
    )

    page_break(document)
    heading(document, "Table of Contents", 1)
    body(
        document,
        "The entries below follow the heading styles of this document. In Microsoft Word, right-click the field and choose Update Field if page numbers need to be refreshed after a later edit.",
        first_line=False,
    )
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.15
    toc_field(paragraph, ' TOC \\o "1-3" \\h \\z \\u ')

    page_break(document)
    heading(document, "List of Tables", 1)
    body(
        document,
        "Tables are numbered by chapter. The captions in this edition are:",
        first_line=False,
    )
    for item in [
        "Table 2.1  Comparison of AICC with nearby alternatives",
        "Table 4.1  User classes and responsibilities",
        "Table 4.2  Functional requirements",
        "Table 4.3  Non-functional requirements",
        "Table 4.4  Use-case catalogue",
        "Table 5.1  Tools and responsibilities",
        "Table 5.2  Indoor route costs",
        "Table 6.1  Subsystem responsibilities",
        "Table 6.2  Design attributes of the main components",
        "Table 7.1  Automated test modules",
        "Table 8.1  Functional test cases",
        "Table 8.2  Chatbot checks",
        "Table 8.3  Navigation checks",
        "Table 8.4  Requirement coverage",
    ]:
        bullet(document, item)

    page_break(document)
    heading(document, "List of Figures", 1)
    body(document, "The figures in Chapter 6 are:", first_line=False)
    for item in [
        "Figure 6.1  Use-case diagram",
        "Figure 6.2  Entity-relationship diagram",
        "Figure 6.3  System architecture",
        "Figure 6.4  Activity diagram for a campus question",
        "Figure 6.5  Sequence diagram for a grounded answer",
        "Figure 6.6  Component diagram",
        "Figure 6.7  Navigation state machine",
        "Figure 6.8  Class diagram of the navigation model",
        "Figure 6.9  Context data-flow diagram",
        "Figure 6.10  Database diagram",
    ]:
        bullet(document, item)

    page_break(document)
    heading(document, "List of Abbreviations", 1)
    body(
        document,
        "Abbreviations are written in full at first use in the body. This list is the reference for the remainder of the report.",
        first_line=False,
    )
    abbreviations = [
        ("AICC", "AI Campus Companion"),
        ("API", "Application Programming Interface"),
        ("BM25", "Best Matching 25, a ranking function for text search"),
        ("CGPA", "Cumulative Grade Point Average"),
        ("CRUD", "Create, Read, Update, and Delete"),
        ("DFD", "Data Flow Diagram"),
        ("ER", "Entity–Relationship"),
        ("FCM", "Firebase Cloud Messaging"),
        ("GPS", "Global Positioning System"),
        ("HTTP", "Hypertext Transfer Protocol"),
        ("JSON", "JavaScript Object Notation"),
        ("LGU", "Lahore Garrison University"),
        ("LLM", "Large Language Model"),
        ("OBE", "Outcome-Based Education"),
        ("QR", "Quick Response code"),
        ("RAG", "Retrieval-Augmented Generation"),
        ("REST", "Representational State Transfer"),
        ("SDK", "Software Development Kit"),
        ("SRS", "Software Requirements Specification"),
        ("TTS", "Text to Speech"),
        ("UI", "User Interface"),
        ("UML", "Unified Modeling Language"),
    ]
    for short, full in abbreviations:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.left_indent = Pt(18)
        run = paragraph.add_run(f"{short}:  {full}")
        _set_run_font(run, size=12)

    page_break(document)
    heading(document, "Abstract", 1)
    body(
        document,
        "Students at Lahore Garrison University already hold several questions at once: where a room is, whether a fee challan is ready, what the attendance rule is, and which office to ask when the rule is unclear. The information exists, but it is split across notices, a handbook, a portal, and the memory of senior students. A general chat model can sound confident while naming a room or a balance that was never stored. A public road map can guide a visitor to the gate and still know nothing about the walkway between that gate and the library.",
    )
    body(
        document,
        "This project presents AICC, a mobile campus companion for that setting. The Android application, built with Flutter, gives a signed-in student a dashboard of stored academic facts, a conversational assistant, teacher profiles, notifications, an outdoor campus map, and indoor floor directions. A FastAPI service holds student records, places, and the handbook. Place questions are answered on the phone from a maintained list, so a coordinate is never invented. Open questions are answered only after handbook passages are retrieved with BM25 and embeddings. A language model on the server may phrase those passages. It may not invent attendance, a CGPA, a document status, or a teacher’s live location. Outdoor walking uses Dijkstra on a campus graph. Indoor walking uses A* on a building graph, with an explicit cost for stairs and for the lift. Outside the campus outline, the application hands the walk to Google Maps.",
    )
    body(
        document,
        "The work was built in short, demonstrable increments and checked with automated tests for place matching, intent, outdoor routing, and indoor routing, plus a scripted set of chatbot and navigation checks. A surveyed GPS-error log is still empty, and the Block B floor drawings used in the demonstration are a sketch until official plans are approved. Within those limits, the system keeps answers attached to records the university can edit, and it gives a student a path the public map does not contain.",
    )
