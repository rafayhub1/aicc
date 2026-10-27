"""Chapters 7 to 9 and references."""

from docx.shared import Cm

from styles import add_table, body, bullet, heading, page_break, table_caption


def add_chapter_7(document):
    page_break(document)
    heading(document, "Chapter 7  Implementation and Testing", 1)
    heading(document, "7.1  Development environment", 2)
    body(
        document,
        "The phone application is a Flutter project targeting Android. The service is a Python application run with Uvicorn. Local demonstration uses a machine on the same network as the phone, or the deployed host whose address the login screen can override and remember. Firebase is configured with the FlutterFire options generated for the mobile project. The campus service reads its own service-account credential from the environment. Neither secret is written into the report or the client.",
    )
    body(
        document,
        "Shared widgets were built before features: primary and secondary buttons, cards, text fields, skeletons, empty states, error states with retry, and status pills. Screens use those widgets so a loading dashboard and a loading map fail in the same way. The theme file holds the navy palette, spacing, and the light and dark themes. Appearance is a preference on the device, not an account setting, so it is available before the first successful read of a student row.",
    )
    heading(document, "7.2  Application structure", 2)
    body(
        document,
        "The process starts in the application entry, which prepares shared preferences and Firebase and then runs the root widget. Named routes cover the splash, a setup screen if Firebase is not ready, sign-in, sign-up, the student shell, the administrator shell, and the indoor map. A session object listens to the authentication user. When the account is an administrator, the student shell leaves and the admin shell opens. That split is the implementation of the two user classes in Table 4.1.",
    )
    body(
        document,
        "The student shell is a tabbed frame: Home, UniBot, History, Notifications, and Teachers. A small campus assistant button can open a new chat from Home. Push messages received in the foreground raise a snack bar. A message that opened the app selects the notification tab. Unread state is refreshed from the inbox service. Sign-out asks for confirmation so a stray tap does not drop the session.",
    )
    heading(document, "7.3  Core student functions", 2)
    body(
        document,
        "Home reads the signed-in student and requests courses for that student’s current semester. Average attendance is the mean of the course percentages, not a second number stored beside them and not a number from the model. The CGPA shown is the overall value on the student row, on a 4.0 scale. Documents of type fee challan and roll-number slip render when their status is ready, using a web view for the file address. A pending document stays a status, not a broken page.",
    )
    body(
        document,
        "UniBot sends text through the classifier in the chat feature. Suggested questions on an empty thread cover the paths a reviewer will try: a challan, a roll slip, attendance, the library, the cafe, a teacher, and the attendance policy. A place hit renders a card with an action to show the map. A remote answer renders the service text and any cards whose identifiers came back with it. Voice capture writes a short recording, uploads it for transcription, and then runs the same classifier on the transcript. Spoken output uses the platform text-to-speech engine and can be turned off. History lists sessions stored for the account, replays a session, and deletes one or all of them.",
    )
    body(
        document,
        "The teacher directory reads faculty profiles and supports search. A profile shows consultation hours so “where is Dr. Hamza Qureshi?” can become a profile rather than a guessed room. Notifications combine the student’s notice list with push. Marking a notice read sends a small update and does not delete the row, so a student can still find it.",
    )
    heading(document, "7.4  Campus guide and indoor map", 2)
    body(
        document,
        "The campus guide loads places and the walkway graph. Search uses the same matcher as chat, so a card and a map pin cannot disagree about which record “canteen” means. Directions call the campus router. The line on the map is the node list. Step text prefers the instruction stored on the edge and otherwise says to continue to the next node’s name. Ending the route clears the session described in Figure 6.7. A QR scan sets the start node from the code an administrator printed for that place, which is how a gate or a stair can be unambiguous when GPS is standing on the wrong side of a wall.",
    )
    body(
        document,
        "The indoor screen loads a building either from a packaged demonstration file or from the remote indoor repository. The floor painter draws the footprint, rooms coloured by type, and the active route. Changing floor does not restart the search. The path already contains the stair or lift link, and the painter shows the part that lies on the selected level. The room sheet names the room under a tap and can become the destination. Preference for stairs is a control on that screen, passed into the router as a cost change rather than as a filter that deletes the lift from the graph. Deleting the lift would hide a real exit. Penalising it keeps it available and rarely chosen.",
    )
    heading(document, "7.5  Service implementation", 2)
    body(
        document,
        "The service process loads environment values, warms handbook embeddings when the file is readable, and mounts the routers under an API prefix. The browser admin surface is a static mount used to inspect demonstration data. Health is a plain status for the milestone that first proved the phone could reach the host. Student writes validate the registration pattern and the department code. Course attendance and CGPA are bounded. Place identifiers are restricted to a machine-readable form so a chat card and a document key cannot drift apart because of a capital letter.",
    )
    body(
        document,
        "Chat accepts a message and a short history. It asks the model for wording and, independently, asks the handbook plan for place identifiers. A failure of the place step does not invent identifiers. It returns an empty list. Transcription rejects a clip that is too short to be speech and a clip that is too large for the endpoint. Groq errors are mapped to a service-unavailable response when the key is rejected or the quota is exhausted, and the phone shows that as an error with retry rather than as a handbook fact.",
    )
    body(
        document,
        "Push registration stores a device token on the student and removes it on sign-out. A broadcast walks students and writes a notification, then attempts delivery. The activity log records administrative actions and can show the actor’s profile when it exists. Statistics for the admin dashboard count the campus record and are loaded beside the faculty count from the mobile store. The dashboard therefore admits, by being two numbers from two stores, that identity and campus records were not forced into one database.",
    )
    heading(document, "7.6  Seed data and content", 2)
    body(
        document,
        "A seed script can create a repeatable set of students across the departments, semesters, and sections in the catalogue, with courses, notices, and documents. Re-running it is safe. A wipe option exists for a clean demonstration and is not the default, so a manual test student is not destroyed by habit. Places and the walkway graph are generated as data files and packaged with the application as well as served remotely, which is why a place answer can still be shown when a reviewer is standing in a poor network spot after the first launch. The handbook used by retrieval is the institutional question set maintained beside the service, with English, Urdu, and Roman Urdu questions. Policy sentences that also exist on the phone for offline answers are kept aligned with that set so the two paths do not contradict.",
    )
    heading(document, "7.7  Testing methodology", 2)
    body(
        document,
        "Testing is layered. Pure functions are unit tested. A small number of widgets are tested for the states a reviewer sees, such as a denied location. The service is tested for the place-link rule and for indoor map validation. A scripted manual pass, written as a demo and as a viva list, covers the sentences that cross the phone and the service together. There is no claim of a full usability trial with a counted cohort. The evaluation that was actually prepared is the one in Chapter 8.",
    )
    body(
        document,
        "Table 7.1 names the automated modules and the requirement they protect. A module is listed because it exists in the repository and asserts a behaviour. This report does not replace that assertion with a decorative pass rate.",
    )
    table_caption(document, "Table 7.1  Automated test modules")
    add_table(
        document,
        ["Module", "What it locks"],
        [
            ["place_search_test", "Longer place names and aliases, including the cafe"],
            ["navigation_intent_test and intent_test", "Place questions versus hours, greetings, and documents"],
            ["campus_router_test", "Dijkstra on the published outdoor graph"],
            ["indoor_router_test", "A* and Dijkstra agreement, stairs and lift"],
            ["indoor_validator_test and demo_building_test", "The demonstration building is a legal graph"],
            ["geo_converter_test", "Local metres and the building origin"],
            ["indoor_map_controller_test and screen tests", "Floor selection and route presentation"],
            ["signup_format_test and api_models_test", "Registration shape and API models"],
            ["academics_test", "Attendance average over the current courses"],
            ["test_place_links", "Cafe card only when the passage or food rule allows it"],
            ["test_groq_prompt", "The prompt carries passages and does not ask for invented marks"],
            ["test_indoor_maps", "Indoor area writes used by the administrator"],
        ],
    )
    heading(document, "7.8  Accuracy, performance, and scale", 2)
    body(
        document,
        "Accuracy is defined here as fidelity to a store, not as a percentage of human preferences. A place answer is accurate when the card is the record that owns that alias. A rule answer is accurate when the numbers match the passage. A route is accurate when it is a minimum path on the published edges. Under that definition the design is accurate by construction for places and routes, and accurate for rules to the extent retrieval returns the right passage. Retrieval can still miss a paraphrase that shares neither words nor embedding neighbourhood. The refusal path is the designed response, and it is preferable to a fluent miss. What is not accurate yet is the outdoor metre error against a surveyed walk. The field log has columns for expected path, actual path, and GPS error, and the measurement cells are empty. Chapter 8 leaves them empty.",
    )
    body(
        document,
        "Performance on the prepared data is interactive because the graphs are small and the place list is a few dozen records. The model call is the slow step, and it is not on the place path or the route path. A cold host on a free tier can add a long first wait. The waking banner is the performance design for that case. It does not make the host faster. It stops the student from treating a sleep as a crash.",
    )
    body(
        document,
        "Scale is data scale. A new place is a document. A new walkway is an edge. A new rule is a handbook block. The routers and the ranker do not contain a case for each room. The limit that will arrive first is not algorithmic. It is the quality of the survey and the discipline of keeping aliases unique. Two places that share the alias “office” would make the longer-name rule ambiguous. The admin form is where that discipline has to live.",
    )
    heading(document, "7.9  Summary", 2)
    body(
        document,
        "The implementation follows the design cut: Flutter for interaction and local decisions, FastAPI for records and grounded wording, Firebase for identity and transcripts, and explicit graphs for both kinds of walk. Tests lock the decisions that are pure. Chapter 8 reports the scripted checks and the limits a reviewer should say aloud.",
    )


def add_chapter_8(document):
    page_break(document)
    heading(document, "Chapter 8  Results and Discussion", 1)
    heading(document, "8.1  Evaluation setup", 2)
    body(
        document,
        "The evaluation is a structured demonstration against the success criteria in Section 3.6 and the use cases in Table 4.4. The phone is signed in as a seeded student and, for maintenance cases, as an administrator. The place list includes the Main Gate, Gate 3, LGU Cafe, Central Library, Admin Block, Masjid, Computer Labs, Classroom B-104, and the other published pins. The outdoor graph connects those pins through walkway nodes such as the south walk, the central paved area, and the admin walk. The indoor demonstration is Block B, with a ground floor and a first floor joined by stairs and a lift. No separate radio map was built, so none is reported.",
    )
    heading(document, "8.2  Functional test cases", 2)
    body(
        document,
        "Table 8.1 is the system test for the use cases. The expected column is the requirement. The observation column states what the implementation does when the case is followed. Where a measurement was not taken, the cell says so.",
    )
    table_caption(document, "Table 8.1  Functional test cases")
    add_table(
        document,
        ["Case", "Use case", "Expected", "Observation"],
        [
            ["TC-01", "UC-01", "Valid sign-in opens Home", "Session binds the registration number and opens the student shell"],
            ["TC-02", "UC-01", "Admin sign-in opens the admin shell", "The student shell redirects when the role is administrator"],
            ["TC-03", "UC-02", "Dashboard shows stored CGPA and courses", "Values come from the student and course payloads"],
            ["TC-04", "UC-03", "A ready challan opens", "Pending status does not pretend to be a file"],
            ["TC-05", "UC-04", "“Where is the library?” shows Central Library", "Matched on the phone, card without waiting for the model"],
            ["TC-06", "UC-08", "Directions on campus draw the graph route", "Dijkstra node list, steps, and End"],
            ["TC-07", "UC-08", "Directions off campus", "A public walking link opens toward a gate"],
            ["TC-08", "UC-08", "Location denied", "Banner, then an “I’m at” list when directions start"],
            ["TC-09", "UC-09", "Block B, first floor, room B-104", "Floor switch and an indoor route from entrance or QR"],
            ["TC-10", "UC-05", "Library hours are not treated as directions", "Hours text; card only from the place tag"],
            ["TC-11", "UC-05", "A question outside the handbook", "Refusal, and no place card"],
            ["TC-12", "UC-13", "Admin drops a pin and saves", "The place appears on the student map after save"],
            ["TC-13", "UC-14", "Broadcast", "A notice is written for students and push is attempted"],
            ["TC-14", "UC-06", "Attendance question", "Course breakdown from the record, not from the model"],
        ],
    )
    heading(document, "8.3  Chatbot checks", 2)
    body(
        document,
        "Table 8.2 is the sentence list used to see that the router and the handbook do not steal each other’s jobs. Checks that are pure place or intent decisions are covered by the phone tests named in Table 7.1. The food question that must attach the cafe from a passage tag is covered by the service test. The refusal case is the one a reviewer should type, because a model that answers a cricket question has left the handbook.",
    )
    table_caption(document, "Table 8.2  Chatbot checks")
    add_table(
        document,
        ["Input", "Path", "Expected"],
        [
            ["LGU", "On-phone place match", "LGU Main Campus card"],
            ["Where is LGU cafe", "On-phone place match", "LGU Cafe, because the longer name wins"],
            ["Where is fountain ground", "On-phone place match", "The fountain place card"],
            ["Navigate to canteen", "Alias", "LGU Cafe"],
            ["Where can I eat lunch?", "Retrieval", "Handbook text plus the cafe card from the place tag"],
            ["When does the library open?", "Retrieval", "Library hours, not a directions session"],
            ["Who won the cricket world cup?", "Retrieval", "Polite refusal, no place card"],
            ["Library kahan hai", "On-phone place match", "Central Library card"],
        ],
    )
    heading(document, "8.4  Navigation checks", 2)
    body(
        document,
        "Table 8.3 lists the walks a demonstration is expected to show. The expected column is the authored graph, not a measured survey. The GPS-error column of the project’s field log is intentionally not copied in as numbers, because that log has not been filled. Publishing blank cells as “within 5 metres” would be a false result.",
    )
    table_caption(document, "Table 8.3  Navigation checks")
    add_table(
        document,
        ["Route", "Expected on the published graph", "Measured GPS error"],
        [
            ["Main Gate to LGU Cafe", "South walk, then toward the cafe", "Not surveyed in this delivery"],
            ["Main Gate to Central Library", "South, central, west, then the library node", "Not surveyed in this delivery"],
            ["Main Gate to the fountain lawn", "South, central, then the fountain node", "Not surveyed in this delivery"],
            ["Main Gate to Classroom B-104", "Outdoor path to Block B, then stairs or lift to the first floor", "Not surveyed in this delivery"],
            ["Any place from off campus", "Public map opens toward the nearest gate", "Not applicable"],
            ["Indoor stair preference", "Lift links carry the extra penalty and stairs are chosen", "Checked by router tests, not by GPS"],
        ],
    )
    heading(document, "8.5  Discussion against the problem", 2)
    body(
        document,
        "The problem in Chapter 3 was the gap between stored facts and a walking student. The checks that pass by construction are the ones that close the gap without a survey: a place card is a record, a personal number is a field, a route is a path on edges an administrator drew, and a question the handbook does not contain does not grow a place card. Those are the results that distinguish this companion from a general chatbot and from a public map, which was the comparison in Table 2.1.",
    )
    body(
        document,
        "The result that is not yet operational is metre-level trust outdoors. Coordinates in the seed are a sketch of the DHA Phase 6 campus. They are good enough to show the interaction and the algorithm. They are not good enough to tell a student, in a binding way, that the cafe is a certain number of metres from the gate. The interface already has the honesty mechanism for a missing edge length: the route is marked estimated. A future survey fills lengths and the mark disappears where the tape measure exists. Until then, a demonstration should say “this is the published path” and should not say “this is a surveyed distance”.",
    )
    body(
        document,
        "Indoor results are similarly split. The search, the floor change, and the instruction builder are real and tested. The picture under the route is a Block B sketch. A student can follow the graph relative to rooms drawn on that sketch. A student should not treat the sketch as an architectural sheet. QR start removes the worst indoor GPS failures from the demo, which is a result about the problem definition as much as about the code. The system stopped claiming a position it could not hold.",
    )
    heading(document, "8.6  Requirement coverage", 2)
    table_caption(document, "Table 8.4  Requirement coverage")
    add_table(
        document,
        ["Requirements", "Evidence", "Residual risk"],
        [
            ["FR-01 to FR-04, FR-25", "Session, dashboard, documents, themed states", "A sleeping host delays the first load"],
            ["FR-05 to FR-12", "Intent tests, chatbot table, voice path", "A paraphrase may miss and must refuse"],
            ["FR-13 to FR-14", "Inbox and faculty screens", "Push can arrive late; the list remains"],
            ["FR-15 to FR-20", "Router tests and navigation table", "Unsurveyed outdoor metres"],
            ["FR-21 to FR-24", "Admin shell and place-save check", "Alias clashes if editors are careless"],
            ["NFR-01 to NFR-03", "Lookups, server key, per-student reads", "Handbook text can still be outdated relative to a new paper notice"],
            ["NFR-04 to NFR-08", "On-phone place and route path, unit tests", "Model quota surfaces as an error, which is intended"],
            ["NFR-09 to NFR-12", "Roman Urdu sample, estimated flag, activity log", "Full Urdu speech was not required"],
        ],
    )
    heading(document, "8.7  Limitations", 2)
    bullet(document, "Outdoor coordinates still need a campus survey before GPS error can be reported.")
    bullet(document, "Floor pictures are a sketch until official drawings are cleared and converted.")
    bullet(document, "Live bus positions, hostel booking, and the registration transaction itself are out of scope.")
    bullet(document, "The handbook can lag a new office notice. The assistant is not the notice.")
    bullet(document, "A free hosting tier sleeps. The first question after a pause waits for the host.")
    bullet(document, "English and Roman Urdu are supported for the prepared phrases. Open Urdu speech was not the acceptance test.")
    heading(document, "8.8  Summary", 2)
    body(
        document,
        "The evaluation shows a companion that answers from records, routes on authored graphs, and refuses what it does not store. It does not show a surveyed positioning error, and this chapter does not invent one. Chapter 9 states what that means for the original problem and what should be done next.",
    )


def add_chapter_9(document):
    page_break(document)
    heading(document, "Chapter 9  Conclusion and Future Work", 1)
    heading(document, "9.1  Conclusion", 2)
    body(
        document,
        "This project set out to give an LGU student one companion for a place, a rule, and a personal document, and to draw a walk on paths the public map does not contain. The delivered system does that by refusing to use one technique for every job. Personal facts are read from the student record. Places are read from a list with aliases. Rules are retrieved from a handbook and only then phrased. Outdoor walks are Dijkstra on published edges. Indoor walks are A* on a building graph with an explicit stair and lift cost. Outside the campus, city walking stays with a public map. The model key stays on the server.",
    )
    body(
        document,
        "The problem statement asked for those separations because a fluent invention is more damaging on a campus than a slow honest answer. The checks in Chapter 8 support the separations that do not need a survey: the library card, the canteen alias, the hours question that is not a route, the cricket question that is refused, the admin pin that appears after save, and the indoor floor change that can prefer stairs. The evaluation that would turn the outdoor sketch into an operational map was prepared as a log and was not filled. The conclusion is therefore precise. The companion works as a grounded campus system on the data it was given. It is not yet a surveyed positioning product.",
    )
    heading(document, "9.2  How the objectives were met", 2)
    bullet(document, "The dashboard objective is met by reading CGPA, courses, attendance, and ready documents from the campus record.")
    bullet(document, "The place objective is met by on-phone matching, including longer names, the cafe alias, and a Roman Urdu sample.")
    bullet(document, "The handbook objective is met by BM25 and embeddings in front of a short, quoted answer, with a stored refusal when nothing matches.")
    bullet(document, "The outdoor objective is met by Dijkstra inside the campus and a public handoff outside it.")
    bullet(document, "The indoor objective is met by A*, QR or chosen start, and stair versus lift costs.")
    bullet(document, "The administrator objective is met by student, teacher, notice, place, and indoor editing without an application update.")
    bullet(document, "The security objective is met by keeping the model credential in the service environment.")
    heading(document, "9.3  Recommendations", 2)
    body(
        document,
        "Before an operational rollout, the walkway nodes should be replaced with points taken on site, and edge lengths should be measured so the estimated flag becomes rare. Official floor drawings should replace the Block B sketch under the same node identifiers, so the router does not have to change when the picture does. Handbook passages should be reviewed by the offices they describe, on a calendar, because a grounded answer can still be an old answer. Alias lists should be unique. Administrators should treat a shared nickname as a defect, not as a convenience.",
    )
    body(
        document,
        "Demonstrations should keep the refusal in the script. A system that only shows its best question teaches the wrong lesson. The cricket question, the pending document, and the off-campus handoff are part of the result.",
    )
    heading(document, "9.4  Future work", 2)
    bullet(document, "A surveyed outdoor graph and a filled GPS-error log, published with the same honesty as the present empty cells.")
    bullet(document, "Replacement of demonstration floor art with cleared architectural plans, without changing the graph contract.")
    bullet(document, "More buildings, each with its own origin, using the same router.")
    bullet(document, "A reviewed Urdu handbook pass so open questions in Urdu script are as well covered as the prepared Roman Urdu place phrases.")
    bullet(document, "Optional inertial smoothing between QR codes, only after the graph itself is trusted, and only as a hint that the student can override.")
    bullet(document, "A tighter loop with the registrar for challan and slip readiness, still as documents the assistant opens rather than as numbers it invents.")
    body(
        document,
        "Live bus tracking remains out of scope until a real vehicle feed exists. Adding a guessed bus position would repeat the failure this project was built to avoid. The next useful bus feature is a stored timetable and a stop that is already a place. A live dot can wait for a live source.",
    )
    body(
        document,
        "The work that should not be undone in that future is the cut itself. Places stay lookups. Rules stay retrieved. Routes stay paths on authored edges. A larger model does not earn the right to skip those steps. If that cut remains, AICC can grow new buildings and new passages without becoming a more confident source of facts the university never stored.",
    )


def add_references(document):
    page_break(document)
    heading(document, "References", 1)
    refs = [
        "[1] E. W. Dijkstra, “A note on two problems in connexion with graphs,” Numerische Mathematik, vol. 1, pp. 269–271, 1959.",
        "[2] P. E. Hart, N. J. Nilsson, and B. Raphael, “A formal basis for the heuristic determination of minimum cost paths,” IEEE Transactions on Systems Science and Cybernetics, vol. 4, no. 2, pp. 100–107, 1968.",
        "[3] P. Lewis et al., “Retrieval-augmented generation for knowledge-intensive NLP tasks,” in Advances in Neural Information Processing Systems, 2020.",
        "[4] S. Robertson and H. Zaragoza, “The probabilistic relevance framework: BM25 and beyond,” Foundations and Trends in Information Retrieval, vol. 3, no. 4, pp. 333–389, 2009.",
        "[5] F. Zafari, A. Gkelias, and K. K. Leung, “A survey of indoor localization systems and technologies,” IEEE Communications Surveys & Tutorials, vol. 21, no. 3, pp. 2568–2599, 2019.",
        "[6] C. D. Manning, P. Raghavan, and H. Schütze, Introduction to Information Retrieval. Cambridge, U.K.: Cambridge University Press, 2008.",
        "[7] I. Sommerville, Software Engineering, 10th ed. Boston, MA: Pearson, 2016.",
        "[8] L. Bass, P. Clements, and R. Kazman, Software Architecture in Practice, 4th ed. Boston, MA: Addison-Wesley, 2021.",
        "[9] ISO/IEC/IEEE 29148:2018, Systems and software engineering — Life cycle processes — Requirements engineering. Geneva, Switzerland: ISO, 2018.",
        "[10] R. S. Pressman and B. R. Maxim, Software Engineering: A Practitioner’s Approach, 9th ed. New York, NY: McGraw-Hill, 2020.",
        "[11] Flutter team, “Flutter documentation,” 2026. [Online]. Available: https://docs.flutter.dev/",
        "[12] S. Ramírez, “FastAPI documentation,” 2026. [Online]. Available: https://fastapi.tiangolo.com/",
        "[13] Google, “Firebase documentation,” 2026. [Online]. Available: https://firebase.google.com/docs",
        "[14] Groq, “Groq API documentation,” 2026. [Online]. Available: https://console.groq.com/docs",
        "[15] Lahore Garrison University, Student information used as the AICC handbook corpus. Lahore, Pakistan: institutional source, maintained with the campus service.",
        "[16] OpenStreetMap contributors, “OpenStreetMap,” 2026. [Online]. Available: https://www.openstreetmap.org/",
    ]
    for ref in refs:
        paragraph = body(document, ref, first_line=False)
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.75)
